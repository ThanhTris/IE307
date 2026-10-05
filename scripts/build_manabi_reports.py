"""Build three reviewable DOCX snapshots from canonical Manabi sources.

Run with the configured document Python runtime (python-docx required).
This writes reports and freshness metadata only, never task status or app code.
"""
from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

from manabi_report_data import MANIFEST, REPORTS, ROOT, UPDATE_HEADINGS, clean, snapshot, text, working_files


def heading_text(value: str) -> str:
    return re.sub(r"[^\w\s]", " ", clean(value)).strip()


def document(title: str, intro: str, source_hash: str):
    doc = Document()
    # The bundled blank template can carry an inherited blue title rule.
    for xml_root in (doc._element, doc.styles.element):
        for border in list(xml_root.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.75)
    sec.header_distance = sec.footer_distance = Inches(.3)
    for style in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        font = doc.styles[style].font
        font.name, font.color.rgb = "Arial", RGBColor(0, 0, 0)
        doc.styles[style].element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    doc.styles["Title"].font.size = Pt(24)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(12)
    for style in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[style].paragraph_format.keep_with_next = True
    doc.add_paragraph(title, "Title")
    doc.add_paragraph("Dự án môn IE307 • Nhóm sáu thành viên • Snapshot tài liệu", "Subtitle")
    doc.add_paragraph(intro)
    doc.add_paragraph("Nguồn: Markdown/spec/task active trong repository. Đề xuất cần chủ dự án review; chưa có app production hoặc số đo hiệu năng đã đạt.")
    doc.core_properties.title = title
    doc.core_properties.author = "Nhóm Manabi"
    doc.core_properties.subject = "Đặc tả và kế hoạch để review"
    doc.core_properties.keywords = "Manabi source SHA256 " + source_hash
    header = sec.header.paragraphs[0]
    header.text = "MANABI  |  TÀI LIỆU PHÁT TRIỂN"
    header.style = doc.styles["Normal"]
    header.runs[0].font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Manabi • Trang ").font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    return doc


def table(doc, headers, rows, widths):
    result = doc.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    for col, width in zip(result.columns, widths):
        col.width = Inches(width)
    for cell, label, width in zip(result.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = label
    repeat = OxmlElement("w:tblHeader")
    result.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = result.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Inches(width)
            cell.text = clean(str(value))
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            element.set(qn("w:" + key), value)
        borders.append(element)
    result._tbl.tblPr.append(borders)
    for index, row in enumerate(result.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side, value in (("top", "80"), ("bottom", "80"), ("left", "100"), ("right", "100")):
                el = OxmlElement("w:" + side)
                el.set(qn("w:w"), value)
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            props.append(margins)
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "F2F2F2")
                props.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                if index == 0:
                    paragraph.paragraph_format.keep_with_next = True
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.bold = index == 0
    return result


def markdown(doc, body: str, *, skip_title=True):
    lines = body.splitlines()
    index, code, buffer = 0, False, []
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith("```"):
            if code and buffer:
                paragraph = doc.add_paragraph("\n".join(buffer))
                for run in paragraph.runs:
                    run.font.name, run.font.size = "Consolas", Pt(8)
                buffer = []
            code = not code
            index += 1
            continue
        if code:
            buffer.append(lines[index])
            index += 1
            continue
        if not line:
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"\|\s*:?-", lines[index + 1]):
            heads = [clean(v) for v in line.strip("|").split("|")]
            rows = []
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([clean(v) for v in lines[index].strip().strip("|").split("|")])
                index += 1
            # Dense source tables become two columns with the remaining cells labeled.
            if len(heads) > 3 and not (len(heads) == 6 and heads[0] == "Thành viên"):
                rows = [[r[0], "\n".join(f"{h}: {v}" for h, v in zip(heads[1:], r[1:]))] for r in rows]
                heads = [heads[0], "Nội dung"]
            widths = ([1.1, 1.05, 1.2, 1.05, 1.2, 1.4] if len(heads) == 6 else
                      [1.15, 5.85] if len(heads) == 2 else [1.15, 2.45, 3.4])
            table(doc, heads, rows, widths)
            continue
        match = re.match(r"^(#{1,6})\s+(.+)", line)
        if match:
            level = len(match.group(1))
            if not (level == 1 and skip_title):
                doc.add_heading(heading_text(match.group(2)), min(max(level - 1, 1), 3))
        elif re.match(r"^[-*] ", line):
            doc.add_paragraph(clean(line[2:]), "List Bullet")
        elif re.match(r"^\d+\. ", line):
            # Keep canonical numbers; Word's shared List Number template otherwise
            # continues numbering across unrelated source sections.
            paragraph = doc.add_paragraph(clean(line))
            paragraph.paragraph_format.left_indent = Pt(14)
            paragraph.paragraph_format.first_line_indent = Pt(-14)
        elif line.startswith("> "):
            doc.add_paragraph(clean(line[2:]))
        else:
            doc.add_paragraph(clean(line))
        index += 1


def source_section(doc, files, path, title, *, page_break=True):
    heading = doc.add_heading(heading_text(title), 1)
    heading.paragraph_format.page_break_before = page_break
    doc.add_paragraph("Nguồn trong repository: " + path)
    markdown(doc, text(files, path))


def sources(doc, files, paths):
    links = {}
    for path in paths:
        for label, url in re.findall(r"\[([^\]]+)\]\((https?://[^)]+)\)", text(files, path)):
            links.setdefault(url, clean(label))
    if not links:
        return
    doc.add_heading("Nguồn tham khảo trực tuyến", 1)
    doc.add_paragraph("Nguồn được liên kết trong tài liệu canonical; URL dưới đây mở được từ Word. Điều khoản/hạn mức cần kiểm tra lại trước triển khai hoặc release.")
    for index, (url, label) in enumerate(links.items(), 1):
        paragraph = doc.add_paragraph(f"{index}. {label}\n")
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), doc.part.relate_to(url, RT.HYPERLINK, is_external=True))
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "17")
        props.append(size)
        run.append(props)
        value = OxmlElement("w:t")
        value.text = url
        run.append(value)
        link.append(run)
        paragraph._p.append(link)


def requirements(files, data):
    doc = document("Manabi yêu cầu và đặc tả chức năng",
                   "Manabi giúp học tiếng Nhật qua flashcard, lịch ôn và ba game. Tài liệu này quy định phạm vi, yêu cầu chức năng và chất lượng, cùng các luồng đủ rõ để người hoặc AI triển khai theo task. AI quiz và ảnh đời sống là pilot có cổng đánh giá.", data["sourceFingerprint"])
    source_section(doc, files, "docs/product/PRODUCT_REQUIREMENTS.md", "Phạm vi và nguyên tắc sản phẩm", page_break=False)
    source_section(doc, files, "docs/product/FUNCTIONAL_REQUIREMENTS.md", "Yêu cầu chức năng")
    source_section(doc, files, "docs/product/NON_FUNCTIONAL_REQUIREMENTS.md", "Yêu cầu phi chức năng")
    names = ("DECK_CARD", "CARD_JSON", "IMPORT", "FLASHCARD", "SRS", "GAMES", "PROGRESS",
             "BACKUP", "AUTH_SYNC", "AI_QUIZ", "IMAGE_CONTEXT", "DATA_PRIVACY", "RELEASE")
    for name in names:
        path = f"docs/specs/{name}_SPEC.md"
        title = text(files, path).splitlines()[0].lstrip("# ")
        source_section(doc, files, path, title)
    sources(doc, files, [f"docs/specs/{name}_SPEC.md" for name in names])
    return doc


def plan(files, data):
    doc = document("Manabi kế hoạch triển khai và lựa chọn dữ liệu",
                   "Nhóm có sáu thành viên và 36 task. Phân công và ước lượng lấy từ nguồn task; tải gồm triển khai, hỗ trợ và review. Core dùng SQLite local và backup JSON; PostgreSQL JSONB thuộc nhánh cloud tùy chọn.", data["sourceFingerprint"])
    for index, (path, title) in enumerate((
        ("docs/project/TEAM_AND_RESPONSIBILITIES.md", "Phân công sáu thành viên"),
        ("docs/project/PROJECT_PLAN.md", "Roadmap và cổng nghiệm thu"),
        ("docs/specs/DATA_STORAGE_SPEC.md", "Thiết kế lưu trữ"),
        ("docs/research/DATABASE_FEASIBILITY.md", "Đánh giá MongoDB và các lựa chọn"),
        ("docs/architecture/decisions/ADR-004-json-storage-and-database.md", "Đề xuất quyết định database"),
        ("docs/project/RISK_REGISTER.md", "Rủi ro và xử lý"),
        ("docs/project/TEAM_WORKFLOW.md", "Quy tắc báo cáo trước push"),
    )):
        source_section(doc, files, path, title, page_break=False)
    doc.add_heading("Danh mục task và nguồn thực thi", 1)
    doc.add_paragraph("Danh mục đầy đủ: tasks/backlog/MASTER_BACKLOG.md. Mỗi task có acceptance criteria, dependency, test plan và evidence. Bản tiến độ DOCX liệt kê 36 task theo từng thành viên; khi task di chuyển thư mục phải đọc file hiện tại, không lấy status cũ trong JSON registry.")
    sources(doc, files, ["docs/research/DATABASE_FEASIBILITY.md", "docs/architecture/decisions/ADR-004-json-storage-and-database.md", "docs/specs/DATA_STORAGE_SPEC.md"])
    return doc


def progress(files, data):
    doc = document("Manabi tiến độ và lỗi của từng thành viên",
                   "Báo cáo này là snapshot task Markdown để trưởng nhóm kiểm tra trước merge. Trạng thái Done chỉ có khi reviewer khác owner chấp thuận; task review chưa được tính hoàn thành. Chưa có task triển khai nào được mặc nhiên đánh dấu đạt chỉ vì đã có kế hoạch.", data["sourceFingerprint"])
    tasks = data["tasks"]
    states = {s: sum(t["status"] == s for t in tasks) for s in ("backlog", "in-progress", "blocked", "review", "done")}
    table(doc, ["Trạng thái", "Số task triển khai"], list(states.items()), [2, 5])
    doc.add_heading("Cách đọc các task chưa bắt đầu", 1)
    doc.add_paragraph("Với mọi dòng backlog bên dưới: đã làm là kế hoạch/spec/task; chưa có code triển khai, file triển khai thay đổi hoặc test chạy. Còn lại là deliverable và toàn bộ acceptance criteria trong task. Blocker kế hoạch là phạm vi/dependency chưa được duyệt; chưa kiểm kỹ thuật nên chưa kết luận không lỗi. Bước tiếp theo là review dependency rồi nhận task. Reviewer chưa duyệt. Khi task bắt đầu, báo cáo thêm đủ bảy mục cập nhật riêng bên dưới bảng của owner.")
    for member in data["members"]:
        name = member["name"]
        heading = doc.add_heading(heading_text("Tiến độ của " + name), 1)
        heading.paragraph_format.page_break_before = True
        owned = [t for t in tasks if t["owner"] == name]
        review = [t for t in tasks if t["reviewer"] == name]
        doc.add_paragraph(f"Owner {len(owned)} task; reviewer {len(review)} task. Điểm được ước lượng theo phạm vi từng task.")
        rows = [[f"{t['id']}\n{t['title']}", f"{t['status']}\n{t['points']} điểm\nReviewer {t['reviewer']}\nPhụ thuộc: {', '.join(t['deps'])}", t["deliverable"]] for t in owned]
        table(doc, ["Task và nội dung", "Trạng thái và review", "Phần cần triển khai hoặc nghiệm thu"], rows, [2.7, 1.7, 2.6])
        doc.add_paragraph("Task phải review: " + ", ".join(t["id"] for t in review))
        for task in owned:
            if task["status"] != "backlog":
                doc.add_heading(heading_text(task["id"] + " Cập nhật thực thi"), 2)
                table(doc, ["Nội dung", "Cập nhật"], [(h, task["updates"][h]) for h in UPDATE_HEADINGS], [1.55, 5.45])
        doc.add_paragraph("Nguồn từng task: " + "; ".join(t["file"] for t in owned))
    return doc


def main():
    files = working_files()
    data, errors = snapshot(files)
    if errors:
        raise SystemExit("\n".join(errors))
    output = ROOT / "docs/deliverables/report"
    output.mkdir(parents=True, exist_ok=True)
    for name, builder in zip(REPORTS[:3], (requirements, plan, progress)):
        doc = builder(files, data)
        doc.save(ROOT / name)
        print(name)
    qa = ROOT / ".work/manabi"
    qa.mkdir(parents=True, exist_ok=True)
    (qa / "snapshot.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest = {"schemaVersion": 1, "project": "Manabi", "generatedAt": datetime.now(timezone.utc).isoformat(),
                "sourceFingerprint": data["sourceFingerprint"], "sourceFileCount": sum(1 for p in files if p not in (*REPORTS, MANIFEST)),
                "artifacts": {name: hashlib.sha256((ROOT / name).read_bytes()).hexdigest() for name in REPORTS if (ROOT / name).exists()},
                "note": "Snapshot only; does not prove app tests or replace independent human review."}
    (ROOT / MANIFEST).write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("Manifest updated; regenerate after any monitored source change, and after workbook generation.")


if __name__ == "__main__":
    main()
